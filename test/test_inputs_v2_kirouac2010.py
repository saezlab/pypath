"""Source contract of the Kirouac 2010 ligand-receptor input."""

import io
from types import SimpleNamespace

import openpyxl
import pytest
from omnipath_build.silver import SilverExtractor
from omnipath_core.biolink import slot_name
from omnipath_core.naming import Namespace
from omnipath_core.source_attributes import PARTICIPANT_ROLE

from pypath.inputs_v2 import kirouac2010 as mod
from pypath.inputs_v2.parsers import kirouac2010 as parser


# Rows of Supplementary Table S12.
TABLE = [
    ('bone morphogenetic protein 2 (BMP2)', 'BMPR 1A/1B/2'),
    ('CCL14 / CCL15 (MIP-5 / MIP-1delta)', 'CCR 1/10'),
    ('Sonic Hedgehog homolog (SHH)', 'PTCH 1/2 / SMO'),
    ('hepatocyte growth factor (HGF)', 'met'),
    ('Follistatin (FST)', '???'),
    ('matrix metalloproteinase 9 (MMP9)', 'xxx'),
    ('secretory leukocyte protease inhibitor (antileukoproteinase)', 'xxx'),
]


def workbook(rows):
    book = openpyxl.Workbook()
    book.active.title = 'TOC'
    sheet = book.create_sheet('S12')
    sheet.append(['Table S12. Ligand-Receptor Mapping', None])
    sheet.append(['LIGAND', 'RECEPTOR(S)'])
    for row in rows:
        sheet.append(list(row))
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    return SimpleNamespace(result=content)


def extract(row):
    record = mod.resource.datasets()['interactions'].mapper(row)
    assert record is not None
    out = SilverExtractor('kirouac2010', 'interactions')
    out.process_record(record, row, 'fixture:1', 1)
    return record, out


@pytest.mark.parametrize(
    'cell,symbols',
    [
        ('adiponectin, (ADIPOQ)', ['ADIPOQ']),
        ('KIT ligand (KITLG / stem cell factor)', ['KITLG']),
        ('hypocretin (orexin) neuropeptide precursor (HCRT)', ['HCRT']),
        (
            'tumor necrosis factor (ligand) superfamily, member 10 (TNFSF10 / TRAIL)',
            ['TNFSF10'],
        ),
        (
            'lymphotoxin alpha (TNF superfamily, member 1 / TNFSF1 / TNF-beta)',
            ['TNFSF1'],
        ),
        (
            'macrophage stimulating 1 ( MST1 / hepatocyte growth factor-like / HLP)',
            ['MST1'],
        ),
        ('transforming growth factor, alpha (TGFA))', ['TGFA']),
        ('CCL2 (MCP-1 / GDCF)', ['CCL2']),
        ('CXCL12 (stromal cell-derived factor 1)', ['CXCL12']),
        ('CCL14 / CCL15 (MIP-5 / MIP-1delta)', ['CCL14', 'CCL15']),
        # The table's own spellings are kept, not corrected.
        ('interferon, alpha 1 (INFA1)', ['INFA1']),
        ('colony stimulating factor 1 (MCSF)', ['MCSF']),
        ('vascular endothelial growth factor D (VEGF-D)', ['VEGF-D']),
        (
            'wingless-type MMTV integration site family, member 7B (WNT7b)',
            ['WNT7b'],
        ),
        ('insulin-like growth factor 2 (IGF1)', ['IGF1']),
    ],
)
def test_ligand_cell_notation(cell, symbols):
    assert parser.ligand_symbols(cell) == symbols


@pytest.mark.parametrize(
    'cell,symbols',
    [
        ('EGFR', ['EGFR']),
        ('ADIPOR 1/2', ['ADIPOR1', 'ADIPOR2']),
        ('BMPR 1A/1B/2', ['BMPR1A', 'BMPR1B', 'BMPR2']),
        ('ACVR 1/1B/2/2B', ['ACVR1', 'ACVR1B', 'ACVR2', 'ACVR2B']),
        ('IL2R A/B/G', ['IL2RA', 'IL2RB', 'IL2RG']),
        ('IL18R 1/AP', ['IL18R1', 'IL18RAP']),
        ('TNFRSF1 A/B', ['TNFRSF1A', 'TNFRSF1B']),
        ('TNFRSF 17/13b', ['TNFRSF17', 'TNFRSF13b']),
        ('IL1R1/2', ['IL1R1', 'IL1R2']),
        ('NOTCH 1-4', ['NOTCH1', 'NOTCH2', 'NOTCH3', 'NOTCH4']),
        ('FGFR1-4', ['FGFR1', 'FGFR2', 'FGFR3', 'FGFR4']),
        ('FZD 1-10', [f'FZD{number}' for number in range(1, 11)]),
        ('PTCH 1/2 / SMO', ['PTCH1', 'PTCH2', 'SMO']),
        ('TNFRSF 21/25 / CCR9', ['TNFRSF21', 'TNFRSF25', 'CCR9']),
        ('IL6R / IL6RST', ['IL6R', 'IL6RST']),
        ('IL8RB (CXCR2)', ['IL8RB']),
        ('Attractin (ATRN)', ['ATRN']),
        # Hyphens inside a symbol are not a numbered series.
        ('LIF-R', ['LIF-R']),
        ('4-1BB', ['4-1BB']),
        ('met', ['met']),
        ('gp130', ['gp130']),
    ],
)
def test_receptor_cell_notation(cell, symbols):
    assert parser.receptor_symbols(cell) == symbols


@pytest.mark.parametrize(
    'cell,designations',
    [
        # The symbol in capitals is the symbol; the other designation is a synonym.
        ('IL8RB (CXCR2)', [('IL8RB', 'CXCR2')]),
        ('BLR1 (CXCR5)', [('BLR1', 'CXCR5')]),
        ('MST1R (RON)', [('MST1R', 'RON')]),
        # The bracketed symbol is the symbol; the name before it is the synonym.
        ('Attractin (ATRN)', [('ATRN', 'Attractin')]),
        ('PTCH 1/2 / SMO', [('PTCH1', None), ('PTCH2', None), ('SMO', None)]),
        ('met', [('met', None)]),
    ],
)
def test_receptor_cell_second_designation(cell, designations):
    assert parser.receptor_designations(cell) == designations


def test_parser_carries_the_second_designation_of_a_receptor_cell():
    rows = list(
        parser.iter_pairs(
            workbook(
                [
                    ('CXCL1 (melanoma growth stimulating activity, alpha / GRO-a)', 'IL8RB (CXCR2)'),
                    ('matrix metalloproteinase 1 (MMP1)', 'xxx'),
                ]
            )
        )
    )
    assert [
        (r['ligand'], r['receptor'], r['receptor_synonym']) for r in rows
    ] == [('CXCL1', 'IL8RB', 'CXCR2'), (None, None, None)]


def test_second_designation_of_a_receptor_is_its_synonym():
    row = {
        'LIGAND': 'CXCL1 (melanoma growth stimulating activity, alpha / GRO-a)',
        'RECEPTOR(S)': 'IL8RB (CXCR2)',
        'ligand': 'CXCL1',
        'receptor': 'IL8RB',
        'receptor_synonym': 'CXCR2',
    }
    record, out = extract(row)
    assert [(i.type, i.value) for i in record.object.identifiers] == [
        (Namespace.GENESYMBOL, 'IL8RB'),
        (Namespace.SYNONYM, 'CXCR2'),
    ]
    # The ligand has none, and the pair's name is unchanged.
    assert [(i.type, i.value) for i in record.subject.identifiers] == [
        (Namespace.GENESYMBOL, 'CXCL1')
    ]
    assert {e.identifier for e in out.entities.values()} == {'CXCL1', 'IL8RB'}
    receptor = out.entities[out.relations[0].object_entity_key]
    assert {(i['ns'], i['id']) for i in receptor.identifiers} == {
        (Namespace.GENESYMBOL, 'IL8RB'),
        (Namespace.SYNONYM, 'CXCR2'),
    }


@pytest.mark.parametrize('cell', ['EGFR and ERBB2', 'IL1RA/B', 'FGFR (1-4)'])
def test_unknown_receptor_notation_is_rejected(cell):
    with pytest.raises(
        ValueError, match='Unknown Kirouac 2010 receptor notation'
    ):
        parser.receptor_symbols(cell)


@pytest.mark.parametrize(
    'cell',
    [
        'epidermal growth factor',
        'secretory leukocyte protease inhibitor (antileukoproteinase)',
    ],
)
def test_ligand_cell_without_a_symbol_is_rejected(cell):
    with pytest.raises(
        ValueError, match='Unknown Kirouac 2010 ligand notation'
    ):
        parser.ligand_symbols(cell)


def test_parser_expands_the_shorthand_into_pairs():
    rows = list(parser.iter_pairs(workbook(TABLE)))
    pairs = [(row['ligand'], row['receptor']) for row in rows]
    assert pairs == [
        ('BMP2', 'BMPR1A'),
        ('BMP2', 'BMPR1B'),
        ('BMP2', 'BMPR2'),
        ('CCL14', 'CCR1'),
        ('CCL14', 'CCR10'),
        ('CCL15', 'CCR1'),
        ('CCL15', 'CCR10'),
        ('SHH', 'PTCH1'),
        ('SHH', 'PTCH2'),
        ('SHH', 'SMO'),
        ('HGF', 'met'),
        # No receptor is given: one row, no pair, whatever the ligand cell holds.
        (None, None),
        (None, None),
        (None, None),
    ]
    # Every row keeps the cells it was read from.
    assert rows[0]['LIGAND'] == 'bone morphogenetic protein 2 (BMP2)'
    assert rows[0]['RECEPTOR(S)'] == 'BMPR 1A/1B/2'
    assert rows[-3]['RECEPTOR(S)'] == '???'


def test_parser_fails_when_the_columns_are_not_the_expected_ones():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'S12'
    sheet.append(['Table S12', None])
    sheet.append(['CYTOKINE', 'RECEPTOR(S)'])
    sheet.append(['epidermal growth factor (EGF)', 'EGFR'])
    content = io.BytesIO()
    book.save(content)
    content.seek(0)
    with pytest.raises(KeyError):
        list(parser.iter_pairs(SimpleNamespace(result=content)))


def test_pair_reaches_silver_extractor_with_the_table_spelling():
    row = {
        'LIGAND': 'hepatocyte growth factor (HGF)',
        'RECEPTOR(S)': 'met',
        'ligand': 'HGF',
        'receptor': 'met',
    }
    record, out = extract(row)
    assert slot_name(record.predicate) == 'interacts_with'
    # The ligand is the subject, the receptor the object.
    assert [(i.type, i.value) for i in record.subject.identifiers] == [
        (Namespace.GENESYMBOL, 'HGF')
    ]
    assert [(i.type, i.value) for i in record.object.identifiers] == [
        (Namespace.GENESYMBOL, 'met')
    ]
    assert [r.predicate for r in out.relations] == ['interacts_with']
    assert {e.entity_type for e in out.entities.values()} == {'protein'}
    assert {e.taxon for e in out.entities.values()} == {'9606'}
    assert {e.identifier for e in out.entities.values()} == {'HGF', 'met'}


def test_no_effect_or_publication_is_invented():
    record, out = extract(
        {
            'LIGAND': 'epidermal growth factor (EGF)',
            'RECEPTOR(S)': 'EGFR',
            'ligand': 'EGF',
            'receptor': 'EGFR',
        }
    )
    assert {slot_name(a.term) for a in record.annotations} == {'in_taxon'}
    # The participants carry their species and, from the table's ligand and
    # receptor columns, their role.
    for entity in (record.subject, record.object):
        assert {
            slot_name(a.term)
            for a in entity.annotations
            if a.term != PARTICIPANT_ROLE
        } == {'in_taxon'}
    assert all(not e.annotations for e in out.entities.values())


def test_ligand_and_receptor_columns_give_the_participant_roles():
    record, out = extract(
        {
            'LIGAND': 'epidermal growth factor (EGF)',
            'RECEPTOR(S)': 'EGFR',
            'ligand': 'EGF',
            'receptor': 'EGFR',
        }
    )
    assert [
        a.value for a in record.subject.annotations if a.term == PARTICIPANT_ROLE
    ] == ['ligand']
    assert [
        a.value for a in record.object.annotations if a.term == PARTICIPANT_ROLE
    ] == ['receptor']
    assert all(a.term != PARTICIPANT_ROLE for a in record.annotations)
    (relation,) = out.relations
    assert sorted(
        (a['scope'], a['value'])
        for a in relation.annotations
        if a['term'] == PARTICIPANT_ROLE
    ) == [('object', 'receptor'), ('subject', 'ligand')]
    assert {e.identifier for e in out.entities.values()} == {'EGF', 'EGFR'}


def test_ligand_without_a_receptor_emits_no_relation():
    row = {
        'LIGAND': 'Follistatin (FST)',
        'RECEPTOR(S)': '???',
        'ligand': None,
        'receptor': None,
    }
    assert mod.map_kirouac_interaction(row) is None
    assert (
        mod.map_kirouac_interaction({'ligand': 'FST', 'receptor': None}) is None
    )

"""Source contract of the Almen2009 input."""

import io
from types import SimpleNamespace

import openpyxl
import pytest
from omnipath_build.silver import SilverExtractor
from omnipath_core.naming import Namespace

from pypath.inputs_v2._almen_classes import CLASS_NAMESPACE
from pypath.inputs_v2.almen2009 import resource
from pypath.inputs_v2.parsers.almen2009 import COLUMNS, class_nodes, iter_classes
from pypath.inputs_v2.parsers.surfaceome import (
    ALMEN_COLUMNS,
    MASTER_SHEET,
    iter_almen_classes,
)

# Rows of the Data sheet of the additional file of Almen et al. 2009: the
# source of the classes.
ROWS = [
    ['IPI00218646', 'CYBB Cytochrome b-245 heavy chain', 'Enzymes', '1', 'Valid', 5.0, 0.0, 'o12-33i45-69o103-121i167-192o204-221i', 5.0, '11-33,50-72,169-191,207-229,401-423', 4.0, 'o10-32i98-120o168-190i202-224o'],
    ['IPI00177543', 'PAM peptidylglycine alpha-amidating monooxygenase isoform a, preproprotein', 'Enzymes', '1.14.17.3/4.3.2.5', 'Valid', 1.0, 'Y', 'n7-14c20/21o863-887i', 1.0, '843-865', 1.0, 'o845-867i'],
    ['IPI00297030', 'TMPRSS11E Transmembrane protease, serine 11E precursor', 'Enzymes', '3.4.', 'Valid', 1.0, 0.0, 'o20-43i', 1.0, '18-40', 1.0, 'i21-43o'],
    ['IPI00011218', 'CSF1R Macrophage colony-stimulating factor 1 receptor precursor', 'Receptors', 'KInase;PDGFR', 'Valid', 1.0, 'Y', 'n4-14c19/20o514-538i', 1.0, '496-518', 1.0, 'o496-518i'],
    ['IPI00002437', 'ADRA2C Alpha-2C adrenergic receptor', 'Receptors', 'GPCR;Rhodopsin', 'Invalid', 7.0, 'Y', 'n5-16c34/35o50-76i88-109o115-147i168-188o208-231i380-401o421-440i', 7.0, '18-40,55-77,91-113,135-157,178-200,347-369,387-409', 7.0, 'o20-42i54-76o91-113i138-160o175-197i345-367o387-406i'],
    ['IPI00443898', 'SLC12A1 Solute carrier family 12 member 1', 'Transporters', 'SLC;APC;SLC12', 'Valid', 10.0, 0.0, 'i196-222o256-280i301-320o326-347i379-398o418-438i545-564o570-591i603-627o913-932i', 11.0, '176-198,212-234,261-283,299-320,327-348,379-401,416-438,487-509,544-565,568-588,612-634', 10.0, 'o182-204i211-233o258-280i301-318o328-350i379-401o416-438i544-561o565-587i607-629o'],
    ['IPI00788970', '- 10 kDa protein', 'Unclassified', 'Unclassified', 'Invalid', 1.0, 0.0, 'o57-78i', 0.0, 0.0, 1.0, 'o57-79i'],
    ['IPI00165934', 'GPR98 Isoform 1 of G-protein coupled receptor 98 precursor', 'Receptors', 'GPCR;Adhesion', 'Valid', 7.0, 'Y', 'n14-25c29/30o5909-5929i5941-5960o5980-6000i6012-6037o6057-6081i6108-6125o6131-6156i', 'ERROR', 'ERROR', 7.0, 'o5880-5902i5911-5933o5948-5970i5982-6004o6030-6052i6073-6095o6105-6127i'],
    ['IPI00394851', 'KLRG2 killer cell lectin-like receptor subfamily G, member 2', 'Receptors', 'SCAR;Class_E', 'Valid', 1.0, 0.0, 'o261-282i', 1.0, '264-286', 1.0, 'o261-283i'],
]
CYBB, PAM, TMPRSS11E, CSF1R, ADRA2C, SLC12A1, UNCLASSIFIED, GPR98, KLRG2_IPI = ROWS

# Rows of the sheet SurfaceomeMasterTable of table S3 of Bausch-Fluck et al.
# 2018, in the order of ALMEN_COLUMNS: entry name, accession, gene, GeneID,
# Almen main class, Almen sub-class. The source of the proteins.
TRBC2 = ['TRBC2_HUMAN', 'A0A5B9', 'TRBC2', None, None, None]
TM129 = ['TM129_HUMAN', 'A0AVI4', 'TMEM129', 92305, 'Unclassified', 'Unclassified']
KLRG2 = ['KLRG2_HUMAN', 'A4D1S0', 'KLRG2', 346689, 'Receptors', 'SCAR;Class_E']
D42E2 = ['D42E2_HUMAN', 'A6NKP2', 'SDR42E2', None, 'Enzymes', 1.1]
S12A8 = ['S12A8_HUMAN', 'A0AV02', 'SLC12A8', 84561, 'Transporters', 'SLC;APC;SLC12']
CY24B = ['CY24B_HUMAN', 'P04839', 'CYBB', 1536, 'Enzymes', 1]
PERT = ['PERT_HUMAN', 'P07202', 'TPO', 7173, 'Enzymes', 1.11]
AMD = ['AMD_HUMAN', 'P19021', 'PAM', 5066, 'Enzymes', '1.14.17.3/4.3.2.5']
MASTER = [TRBC2, TM129, KLRG2, D42E2, S12A8, CY24B, PERT, AMD]


def workbook(rows, sheet='Data', header=COLUMNS):
    book = openpyxl.Workbook()
    book.active.title = 'Legend'
    data = book.create_sheet(sheet)
    data.append(list(header))
    for row in rows:
        data.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleNamespace(result=io.BytesIO(buffer.getvalue()))


def master(rows, header=ALMEN_COLUMNS, sheet=MASTER_SHEET):
    book = openpyxl.Workbook()
    table = book.active
    table.title = sheet
    table.append(['table S3: List of human proteins with results from surfaceome classification'])
    table.append(list(header))
    for row in rows:
        table.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleNamespace(result=io.BytesIO(buffer.getvalue()))


def records(rows):
    return list(iter_almen_classes(master(rows)))


def extract(row, dataset='proteins'):
    record = resource.datasets()[dataset].mapper(row)
    assert record is not None
    out = SilverExtractor('almen2009', dataset)
    out.process_record(record, row, 'fixture:1', 1)
    return out


def relations_of(row):
    out = extract(row)
    return [
        (out.entities[r.subject_entity_key], r, out.entities[r.object_entity_key])
        for r in out.relations
    ]


def relations_of_all(parsed):
    return [triple for record in parsed for triple in relations_of(record)]


def attributes(relation):
    return {a['term']: a['value'] for a in relation.annotations if a['scope'] == 'relation'}


def test_the_proteins_are_read_from_the_surfaceome_table_and_the_classes_from_the_paper():
    datasets = resource.datasets()
    assert list(datasets) == ['terms', 'proteins']
    assert datasets['terms'].download.filename == 'almen2009_12915_2009_258_MOESM1_ESM.xls'
    assert datasets['proteins'].download.filename == 'table_S3_surfaceome.xlsx'


def test_parser_reads_one_statement_per_class_of_a_protein():
    by_accession = {}
    for record in records(MASTER):
        by_accession.setdefault(record['uniprot'], []).append((record['kind'], record['term']))
    assert by_accession == {
        'A0AVI4': [('class', 'Unclassified')],
        'A4D1S0': [('class', 'SCAR;Class_E')],
        # 1.1 as a number is EC 1.1 or EC 1.10: the main class only.
        'A6NKP2': [('class', 'Enzymes')],
        'A0AV02': [('class', 'SLC;APC;SLC12')],
        # Numbers that stand for one EC class.
        'P04839': [('class', 'Enzymes'), ('ec', '1')],
        'P07202': [('class', 'Enzymes'), ('ec', '1.11')],
        'P19021': [('class', 'Enzymes'), ('ec', '1.14.17.3'), ('ec', '4.3.2.5')],
    }
    # A protein the sheet gives no class to is not in the classification.
    assert 'A0A5B9' not in by_accession


def test_parser_keeps_the_identifiers_of_the_sheet():
    by_accession = {r['uniprot']: r for r in records(MASTER)}
    assert by_accession['A0AV02'] == {
        'uniprot': 'A0AV02',
        'entry_name': 'S12A8_HUMAN',
        'gene': 'SLC12A8',
        'entrez': ['84561'],
        'kind': 'class',
        'term': 'SLC;APC;SLC12',
    }
    assert by_accession['A6NKP2']['entrez'] == []


def test_an_ec_class_written_as_text_is_read_as_written():
    # Only a number can have lost a zero: the text 1.1 is EC 1.1.
    row = [*D42E2[:5], '1.1']
    assert [(r['kind'], r['term']) for r in records([row])] == [('class', 'Enzymes'), ('ec', '1.1')]


@pytest.mark.parametrize(
    'change',
    [
        {4: 'Plants'},  # unknown main class
        {4: 'Enzymes', 5: 'GPCR;Rhodopsin'},  # an enzyme without an EC number
        {5: 'GPCR;'},  # an empty class
        {4: None},  # a sub-class without a main class
        {5: 7},  # a class of a receptor that is a number
        {3: 'x'},  # not an Entrez Gene ID
        {1: None},  # no accession
    ],
)
def test_parser_raises_on_a_value_it_does_not_know(change):
    row = list(KLRG2)
    for index, value in change.items():
        row[index] = value
    with pytest.raises(ValueError):
        records([row])


def test_parser_raises_on_a_repeated_accession_missing_columns_and_no_class():
    with pytest.raises(ValueError, match='occurs twice'):
        records([KLRG2, KLRG2])
    with pytest.raises(ValueError, match='lacks expected columns'):
        list(iter_almen_classes(master([KLRG2], header=ALMEN_COLUMNS[:-1])))
    with pytest.raises(ValueError):
        records([])
    # A sheet where no protein has a class of Almen et al.
    with pytest.raises(ValueError, match='gives no protein a class'):
        records([TRBC2])


def test_classes_parser_checks_the_additional_file():
    with pytest.raises(ValueError, match='occurs twice'):
        list(iter_classes(workbook([CYBB, CYBB])))
    with pytest.raises(ValueError, match='Unknown main class'):
        list(iter_classes(workbook([[*ADRA2C[:2], 'Plants', *ADRA2C[3:]]])))
    with pytest.raises(ValueError, match='Not an EC number'):
        list(iter_classes(workbook([[*ADRA2C[:2], 'Enzymes', *ADRA2C[3:]]])))
    with pytest.raises(ValueError, match='lacks expected columns'):
        list(iter_classes(workbook([CYBB], header=COLUMNS[:-1])))
    with pytest.raises(ValueError):
        list(iter_classes(workbook([])))


def test_a_class_path_has_one_main_class():
    with pytest.raises(ValueError, match='two main classes'):
        list(class_nodes([('Receptors', 'Kinase;Other'), ('Miscellaneous', 'Kinase;Other')]))


def test_classes_are_terms_with_names_and_parents():
    nodes = {n['id']: n for n in iter_classes(workbook(ROWS))}
    assert nodes['Receptors'] == {
        'id': 'Receptors', 'name': 'Receptors', 'label': 'Receptors', 'parent': None
    }
    assert nodes['GPCR'] == {
        'id': 'GPCR', 'name': 'GPCR', 'label': 'GPCR', 'parent': 'Receptors'
    }
    assert nodes['SLC;APC;SLC12'] == {
        'id': 'SLC;APC;SLC12',
        'name': 'SLC > APC > SLC12',
        'label': 'SLC12',
        'parent': 'SLC;APC',
    }
    assert nodes['SLC;APC']['parent'] == 'SLC'
    assert nodes['SLC']['parent'] == 'Transporters'
    # The five main classes are always terms; the sheet's Unclassified class
    # is the main class itself, not a class below it.
    assert {'Enzymes', 'Receptors', 'Transporters', 'Miscellaneous', 'Unclassified'} <= set(nodes)
    assert 'Kinase;PDGFR' in nodes and 'KInase;PDGFR' not in nodes


def test_the_vocabulary_dataset_is_called_terms():
    assert 'terms' in resource.datasets()
    assert 'classes' not in resource.datasets()


def test_class_term_has_name_synonym_and_subclass_of_its_parent():
    row = {'id': 'SLC;APC;SLC12', 'name': 'SLC > APC > SLC12', 'label': 'SLC12', 'parent': 'SLC;APC'}
    out = extract(row, 'terms')
    (term,) = [e for e in out.entities.values() if len(e.identifiers) == 3]
    assert (term.entity_type, term.namespace, term.identifier) == (
        'ontology_class', CLASS_NAMESPACE, 'SLC;APC;SLC12'
    )
    assert {(i['ns'], i['id']) for i in term.identifiers} == {
        (CLASS_NAMESPACE, 'SLC;APC;SLC12'),
        (Namespace.NAME, 'SLC > APC > SLC12'),
        (Namespace.SYNONYM, 'SLC12'),
    }
    (relation,) = out.relations
    assert relation.predicate == 'subclass_of'
    parent = out.entities[relation.object_entity_key]
    assert (parent.entity_type, parent.namespace, parent.identifier) == (
        'ontology_class', CLASS_NAMESPACE, 'SLC;APC'
    )


def test_protein_is_associated_with_the_most_specific_class():
    ((protein, relation, term),) = relations_of(records([S12A8])[0])
    assert relation.predicate == 'associated_with'
    assert (protein.entity_type, protein.namespace, protein.identifier, protein.taxon) == (
        'protein', 'uniprot', 'A0AV02', '9606'
    )
    assert {(i['ns'], i['id']) for i in protein.identifiers} == {
        (Namespace.UNIPROT, 'A0AV02'),
        (Namespace.UNIPROT_ENTRY, 'S12A8_HUMAN'),
        (Namespace.GENESYMBOL, 'SLC12A8'),
        (Namespace.ENTREZ, '84561'),
    }
    assert (term.entity_type, term.namespace, term.identifier) == (
        'ontology_class', CLASS_NAMESPACE, 'SLC;APC;SLC12'
    )
    # The sheet has nothing of Almen et al. on the protein beyond its class.
    assert attributes(relation) == {}


def test_object_and_class_term_have_the_same_type_namespace_and_identifier():
    # The classes of the additional file name the objects of the statements
    # read from the surfaceome table.
    nodes = {n['id']: n for n in iter_classes(workbook(ROWS))}
    checked = 0
    for record in records(MASTER):
        if record['kind'] != 'class':
            continue
        ((_, _, obj),) = relations_of(record)
        out = extract(nodes[record['term']], 'terms')
        (defined,) = [
            e for e in out.entities.values()
            if e.identifier == record['term'] and (len(e.identifiers) > 1 or e.label)
        ]
        assert (defined.entity_type, defined.namespace, defined.identifier) == (
            obj.entity_type, obj.namespace, obj.identifier
        )
        checked += 1
    assert checked == 7


def test_an_enzyme_is_in_the_main_class_and_in_its_ec_classes():
    got = {
        (term.namespace, term.identifier)
        for _, _, term in relations_of_all(records([AMD]))
    }
    assert got == {
        (CLASS_NAMESPACE, 'Enzymes'),
        ('ec', 'EC:1.14.17.3'),
        ('ec', 'EC:4.3.2.5'),
    }


def test_an_enzyme_with_an_ambiguous_number_gets_no_ec_class():
    got = [
        (term.namespace, term.identifier)
        for _, _, term in relations_of_all(records([D42E2]))
    ]
    assert got == [(CLASS_NAMESPACE, 'Enzymes')]


def test_nothing_is_asserted_beyond_a_class_for_a_protein():
    # No IPI accession, no predicate other than the association, no attribute.
    for protein, relation, _ in relations_of_all(records(MASTER)):
        assert relation.predicate == 'associated_with'
        assert Namespace.IPI not in {i['ns'] for i in protein.identifiers}
        assert attributes(relation) == {}


def test_an_unknown_kind_of_class_raises():
    row = dict(records([CY24B])[0], kind='family')
    with pytest.raises(ValueError, match='Unknown kind'):
        resource.proteins.mapper(row)

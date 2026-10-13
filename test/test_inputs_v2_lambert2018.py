"""Lambert 2018: reader, mappers and extractor on rows of Table S1 (The Human Transcription Factors)."""

import io
from types import SimpleNamespace

import openpyxl
import pytest
from omnipath_build.silver import SilverExtractor
from omnipath_core.naming import Namespace

from pypath.inputs_v2 import lambert2018
from pypath.inputs_v2.parsers import lambert2018 as parser

GROUPS = [
    'Gene Information', None, None, 'Is TF?', 'Final Assesment', None, None, None, None,
    'Notes from re-reviewed genes', None, None, None, None, 'Curator Disagreement Flags',
    None, 'Curator 1', None, None, None, None, 'Curator 2', None, None, None, None,
    'Prior classifications', None, None, None, None, None, None, None, None, None, None,
]
COLUMNS = [
    'ID', 'Name', 'DBD', None, 'TF assessment', 'Binding mode', 'Motif status', 'Notes',
    'Comments', 'Committee notes', 'MTW Notes', 'TRH Notes', 'SL notes', 'AJ notes',
    'Disagree on Assessment ', 'Disagree on Binding', 'Author1', 'Assesment1', 'Binding1',
    'Comment1', 'Notes1', 'Author2', 'Assesment2', 'Binding2', 'Comment2', 'Notes2',
    'Vaquerizas 2009 TF classification', 'CisBP considers it as a TF?',
    'TFclass considers it as a TF? ', 'TF-CAT classification', 'Is a GO TF', 'PDB',
    None, None, None, None, None,
]


def gene(id_, name, dbd, verdict, assessment, binding, motif, notes=None,
         disagree=(None, None), one=('', ''), two=('', '')):
    row = [None] * 37
    row[:8] = [id_, name, dbd, verdict, assessment, binding, motif, notes]
    row[14:16] = disagree
    row[17], row[18] = one
    row[22], row[23] = two
    return row


MONOMER = '1 Monomer or homomultimer'
LOW = '3 Low specificity DNA-binding protein'
# Rows of Table S1: a known-motif TF, two with the same family, the gene whose
# ID is a UniProt entry name, a TF of unknown family, a gene with the verdict
# No and the pseudogene with a UniProt entry name and the verdict No.
ROWS = [
    gene('ENSG00000137203', 'TFAP2A', 'AP-2', 'Yes', 'Known motif', MONOMER,
         'High-throughput in vitro', None, one=('Has known motif', MONOMER),
         two=('Has known motif', MONOMER)),
    gene('ENSG00000143190', 'POU2F1', 'Homeodomain; POU', 'Yes', 'Known motif', MONOMER,
         'High-throughput in vitro'),
    gene('ENSG00000185668', 'POU3F1', 'Homeodomain; POU', 'Yes', 'Known motif', MONOMER,
         'High-throughput in vitro'),
    gene('DUX1_HUMAN', 'DUX1', 'Homeodomain', 'Yes', 'Known motif', MONOMER,
         'In vivo/Misc source',
         'Only known motifs are from Transfac or HocoMoco - origin is uncertain',
         disagree=('Disagree', None)),
    gene('ENSG00000119403', 'PHF19 ', 'Unknown', 'Yes', 'Likely to be sequence specific TF',
         MONOMER, 'No motif'),
    gene('ENSG00000117713', 'ARID1A', 'ARID/BRIGHT', 'No', 'Unlikely to be sequence specific TF',
         LOW, 'No motif'),
    gene('ZNF73_HUMAN', 'ZNF73', 'C2H2 ZF', 'No', 'Unlikely to be sequence specific TF',
         '4 Not a DNA binding protein', 'No motif', 'Pseudogene'),
]


def workbook(rows=ROWS, second=COLUMNS, first=GROUPS):
    book = openpyxl.Workbook()
    book.create_sheet('README').append(['Table S1. Related to Figure 1B.'])
    sheet = book.create_sheet(parser.SHEET)
    for row in [first, second, *rows]:
        sheet.append(row)
    del book['Sheet']
    data = io.BytesIO()
    book.save(data)
    return SimpleNamespace(result=io.BytesIO(data.getvalue()))


def extract(dataset, rows):
    out = SilverExtractor('lambert2018', dataset)
    mapper = lambert2018.resource.datasets()[dataset].mapper
    for index, row in enumerate(rows):
        record = mapper(row)
        if record is not None:
            out.process_record(record, row, f'fixture:{index}', index)
    return out


def test_reader_checks_both_header_rows_and_the_ids():
    genes = list(parser.iter_transcription_factors(workbook()))
    assert len(genes) == 7
    assert genes[0] == {
        'id': 'ENSG00000137203', 'symbol': 'TFAP2A', 'dbd': 'AP-2', 'is_tf': 'Yes',
        'assessment': 'Known motif', 'binding_mode': MONOMER,
        'motif_status': 'High-throughput in vitro', 'notes': None,
        'assessment_disagreement': None, 'binding_disagreement': None,
        'assessment_1': 'Has known motif', 'assessment_2': 'Has known motif',
        'binding_1': MONOMER, 'binding_2': MONOMER, 'id_type': 'ensg',
    }
    # Cells are stripped ("PHF19 "), and the source's UniProt entry name is kept.
    assert genes[4]['symbol'] == 'PHF19'
    assert (genes[3]['id'], genes[3]['id_type']) == ('DUX1_HUMAN', 'uniprot_entry')
    with pytest.raises(ValueError, match="column 'dbd'"):
        list(parser.iter_transcription_factors(workbook(second=[*COLUMNS[:2], 'Domain', *COLUMNS[3:]])))
    with pytest.raises(ValueError, match="column 'is_tf'"):
        list(parser.iter_transcription_factors(workbook(first=[*GROUPS[:3], 'TF?', *GROUPS[4:]])))
    with pytest.raises(ValueError, match='Unknown Lambert 2018 verdict'):
        list(parser.iter_transcription_factors(workbook(rows=[gene('ENSG00000137203', 'X', 'AP-2', 'Maybe', 'a', 'b', 'c')])))
    with pytest.raises(ValueError, match='Malformed Lambert 2018 gene ID'):
        list(parser.iter_transcription_factors(workbook(rows=[gene('TFAP2A', 'X', 'AP-2', 'Yes', 'a', 'b', 'c')])))
    with pytest.raises(ValueError, match='has no genes'):
        list(parser.iter_transcription_factors(workbook(rows=[])))


def test_genes_with_the_verdict_no_are_not_members_of_the_list():
    genes = list(parser.iter_transcription_factors(workbook()))
    assert lambert2018.map_transcription_factor(genes[5]) is None
    out = extract('transcription_factors', genes)
    assert len(out.relations) == 5
    identifiers = {i['id'] for e in out.entities.values() for i in e.identifiers}
    assert 'ENSG00000117713' not in identifiers and 'ZNF73_HUMAN' not in identifiers
    assert 'ARID1A' not in identifiers and 'ZNF73' not in identifiers


def test_statement_carries_the_assessment_and_the_attributes():
    out = extract('transcription_factors', parser.iter_transcription_factors(workbook()))
    by_gene = {}
    for r in out.relations:
        subject = out.entities[r.subject_entity_key]
        symbol = next(i['id'] for i in subject.identifiers if i['ns'] == Namespace.GENESYMBOL)
        assert r.predicate == 'associated_with'
        by_gene[symbol] = [(a['term'], a['value']) for a in r.annotations if a['scope'] == 'relation']
    assert by_gene['TFAP2A'] == [
        ('has_confidence_level', 'Known motif'),
        ('lambert2018:binding_mode', MONOMER),
        ('lambert2018:motif_status', 'High-throughput in vitro'),
        ('lambert2018:curator_assessment', 'Has known motif'),
        ('lambert2018:curator_assessment', 'Has known motif'),
        ('lambert2018:curator_binding_mode', MONOMER),
        ('lambert2018:curator_binding_mode', MONOMER),
    ]
    dux1 = dict(by_gene['DUX1'])
    assert dux1['lambert2018:assessment_disagreement'] == 'Disagree'
    assert dux1['description'].startswith('Only known motifs')
    # The ID that is a UniProt entry name stays as the source writes it.
    dux = next(
        e for e in out.entities.values()
        if e.entity_type == 'gene' and (Namespace.UNIPROT_ENTRY, 'DUX1_HUMAN') in {(i['ns'], i['id']) for i in e.identifiers}
    )
    assert (Namespace.ENSG, 'DUX1_HUMAN') not in {(i['ns'], i['id']) for i in dux.identifiers}
    assert {e.taxon for e in out.entities.values() if e.entity_type == 'gene'} == {'9606'}
    with pytest.raises(ValueError, match='Unknown Lambert 2018 assessment'):
        lambert2018.map_transcription_factor(
            dict(next(iter(parser.iter_transcription_factors(workbook()))), assessment='Unlikely to be sequence specific TF')
        )


def test_the_object_of_a_statement_is_the_term_of_the_list():
    terms = extract('terms', parser.iter_list(workbook()))
    (term,) = terms.entities.values()
    out = extract('transcription_factors', parser.iter_transcription_factors(workbook()))
    assert term.entity_type == 'ontology_class'
    assert (term.namespace, term.identifier) == ('lambert2018_list', 'human_transcription_factors')
    assert term.label == 'Human transcription factors (Lambert et al. 2018)'
    description = next(a['value'] for a in term.annotations if a['term'] == 'description')
    assert '5 of the 7 candidates' in description
    objects = {
        (e.entity_type, e.namespace, e.identifier)
        for e in (out.entities[r.object_entity_key] for r in out.relations)
    }
    assert objects == {(term.entity_type, term.namespace, term.identifier)}


def test_dbd_families_have_the_tf_genes_as_members():
    families = list(parser.iter_dbd_families(workbook()))
    # Unknown is no family, and a gene with the verdict No is in none.
    assert [(f['dbd'], [m['symbol'] for m in f['members']]) for f in families] == [
        ('AP-2', ['TFAP2A']),
        # A gene with two kinds of domain ("Homeodomain; POU") is in both families.
        ('Homeodomain', ['POU2F1', 'POU3F1', 'DUX1']),
        ('POU', ['POU2F1', 'POU3F1']),
    ]
    out = extract('families', families)
    assert {e.namespace for e in out.entities.values() if e.entity_type == 'gene_family'} == {
        'lambert2018_dbd_family'
    }
    assert {e.taxon for e in out.entities.values()} == {'9606'}
    assert [r.predicate for r in out.relations] == ['has_member'] * 6
    labels = {e.label for e in out.entities.values() if e.entity_type == 'gene_family'}
    assert labels == {'AP-2', 'Homeodomain', 'POU'}


def test_the_vocabulary_dataset_is_called_terms():
    assert list(lambert2018.resource.datasets()) == ['terms', 'transcription_factors', 'families']
